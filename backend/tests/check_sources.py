"""运行：python -m unittest backend.tests.check_sources -v；不访问网络或正式数据库。"""

from io import BytesIO
import os
from pathlib import Path
import subprocess
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

import httpx
from fastapi import HTTPException
from openpyxl import Workbook

from backend.app import import_sources as sources


def xlsx(rows, extra=None):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    if extra is not None:
        sheet = workbook.create_sheet('第二表')
        for row in extra:
            sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class XlsxChecks(unittest.TestCase):
    def rejected(self, data, message):
        with self.assertRaises(HTTPException) as caught:
            sources.parse_xlsx(data)
        self.assertEqual(caught.exception.status_code, 422)
        self.assertIn(message, caught.exception.detail)

    def test_valid_and_blank_rows(self):
        result = sources.parse_xlsx(xlsx([
            ['id', 'words', 'meaning', 'example', 'example_translation'],
            [123, ' apple ', '苹果', 'An apple.', '一个苹果。'], [],
        ], extra=[]))
        self.assertEqual(result, {'sheets': [{'name': 'Sheet', 'items': [{
            'word': 'apple', 'meaning': '苹果', 'example': 'An apple.', 'example_translation': '一个苹果。',
        }]}], 'skipped': [{'name': '第二表', 'reason': '空表或没有可导入的单词'}]})

    def test_invalid(self):
        for rows, message in [
            ([['word', 'words', 'meaning']], '重复表头'),
            ([['word', 'word', 'meaning']], '重复表头'),
            ([['word']], '缺少表头'),
            ([[], ['word', 'meaning']], '第一行'),
            ([['word', 'meaning'], [1, '一']], '必须是文本'),
            ([['word', 'meaning'], ['a', True]], '必须是文本'),
            ([['word', 'meaning'], ['a', None]], 'meaning 不能为空'),
            ([['word', 'meaning', 'example'], ['a', '一', 12]], 'example 必须是文本'),
            ([['word', 'meaning'], ['#DIV/0!', '一']], 'Excel 错误值'),
        ]:
            with self.subTest(rows=rows):
                self.rejected(xlsx(rows), message)
        self.rejected(b'not a zip', 'XLSX 无效')
        self.rejected('not bytes', 'bytes')
        self.rejected(xlsx([]), '没有')
        result = sources.parse_xlsx(xlsx([['word', 'meaning'], ['a', '一']], extra=[['oops']]))
        self.assertEqual(len(result['sheets']), 1)
        self.assertEqual(result['skipped'][0]['name'], '第二表')
        self.assertIn('缺少表头', result['skipped'][0]['reason'])

    def test_multiple_valid_sheets_ignore_readme(self):
        workbook = Workbook()
        workbook.active.title = 'README'
        workbook.active.append(['说明'])
        for name in ('第一组', '第二组'):
            sheet = workbook.create_sheet(name)
            sheet.append(['word', 'meaning'])
            sheet.append(['hello', '你好'])
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        result = sources.parse_xlsx(stream.getvalue())
        self.assertEqual([s['name'] for s in result['sheets']], ['第一组', '第二组'])
        self.assertEqual(result['skipped'][0]['name'], 'README')

    def test_limits_and_forged_dimensions(self):
        self.rejected(b'x' * (sources.MAX_BYTES + 1), '2 MiB')
        self.rejected(xlsx([['word', 'meaning']] + [['a', '一']] * 5001), '5001')
        self.rejected(xlsx([['word', 'meaning'] + [''] * 99]), '100')
        stream = BytesIO()
        with ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('bomb', b'x' * (sources.MAX_EXPANDED + 1))
        self.rejected(stream.getvalue(), '20 MiB')
        stream = BytesIO()
        with ZipFile(stream, 'w') as archive:
            for number in range(2001):
                archive.writestr(str(number), b'')
        self.rejected(stream.getvalue(), '2000')
        stream = BytesIO()
        with ZipFile(BytesIO(xlsx([['word', 'meaning'], ['a', '一']]))) as original:
            with ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
                for entry in original.infolist():
                    body = original.read(entry)
                    if entry.filename == 'xl/worksheets/sheet1.xml':
                        body = body.replace(b'A1:B2', b'A1:A1')
                    archive.writestr(entry.filename, body)
        self.assertEqual(len(sources.parse_xlsx(stream.getvalue())['sheets'][0]['items']), 1)


class DriveChecks(unittest.IsolatedAsyncioTestCase):
    async def fetch(self, url, handler, private=False):
        real_client = httpx.AsyncClient
        with patch.object(sources.httpx, 'AsyncClient', side_effect=lambda **kw: real_client(
            transport=httpx.MockTransport(handler), **kw
        )):
            return await sources.read_drive(url, allow_private=private)

    async def test_fixed_urls_and_redirects(self):
        seen = []

        def handler(request):
            seen.append(str(request.url))
            if len(seen) == 1:
                return httpx.Response(302, headers={'location': 'https://drive.usercontent.google.com/download?id=abc'})
            return httpx.Response(200, content=b'xlsx bytes')

        result = await self.fetch('https://docs.google.com/spreadsheets/d/abc/edit?evil=x', handler)
        self.assertEqual(result, b'xlsx bytes')
        self.assertEqual(seen[0], 'https://docs.google.com/spreadsheets/d/abc/export?format=xlsx')
        for url in ['https://drive.google.com/file/d/abc/view', 'https://drive.google.com/open?id=abc']:
            self.assertEqual(await self.fetch(url, lambda request: httpx.Response(200, content=b'ok')), b'ok')

    async def test_sheets_download_subdomain(self):
        seen = []
        def handler(request):
            seen.append(str(request.url))
            if len(seen) == 1:
                return httpx.Response(307, headers={'location': 'https://doc-0g-6c-sheets.googleusercontent.com/export/file?format=xlsx'})
            return httpx.Response(200, content=b'xlsx')
        self.assertEqual(await self.fetch('https://docs.google.com/spreadsheets/d/abc/edit?gid=0#gid=0', handler), b'xlsx')
        self.assertEqual(len(seen), 2)
        for host in ['doc-a-sheets.googleusercontent.com.evil.com', 'evil.googleusercontent.com', 'user@doc-a-sheets.googleusercontent.com']:
            with self.assertRaises(HTTPException):
                await self.fetch('https://docs.google.com/spreadsheets/d/abc/edit', lambda r: httpx.Response(307, headers={'location': 'https://' + host + '/export'}))

    async def test_reject_urls(self):
        for url in [
            'http://docs.google.com/spreadsheets/d/a', 'https://evil.com/file/d/a',
            'https://docs.google.com.evil.com/spreadsheets/d/a',
            'https://user@docs.google.com/spreadsheets/d/a',
            'https://docs.google.com:443/spreadsheets/d/a',
            'https://drive.google.com/open?id=a&id=b',
            'https://drive.google.com/file/d/a/../../evil',
        ]:
            with self.subTest(url=url), self.assertRaises(HTTPException):
                await sources.read_drive(url)
        with patch.object(sources, '_read_private') as private, self.assertRaises(HTTPException):
            await self.fetch('https://drive.google.com/open?id=a', lambda request: httpx.Response(
                302, headers={'location': 'https://evil.com/secret'}
            ), private=True)
        private.assert_not_called()

    async def test_fallback_gate_and_format(self):
        url = 'https://docs.google.com/spreadsheets/d/a/edit'
        for enabled, allowed in [('0', True), ('1', False), ('1', True)]:
            with patch.dict(os.environ, {'WORDS_GOG_PRIVATE': enabled}), patch.object(
                sources, '_read_private', return_value=b'private'
            ) as private:
                if enabled == '1' and allowed:
                    self.assertEqual(await self.fetch(url, lambda r: httpx.Response(403), allowed), b'private')
                    private.assert_called_once_with('a', True)
                else:
                    with self.assertRaises(HTTPException):
                        await self.fetch(url, lambda r: httpx.Response(403), allowed)
                    private.assert_not_called()
        with patch.dict(os.environ, {'WORDS_GOG_PRIVATE': '1'}), patch.object(sources, '_read_private') as private:
            self.assertEqual(await self.fetch(url, lambda r: httpx.Response(200, content=b'invalid xlsx'), True), b'invalid xlsx')
            with self.assertRaises(HTTPException):
                sources.parse_xlsx(b'invalid xlsx')
            private.assert_not_called()
            with self.assertRaises(HTTPException):
                await self.fetch(url, lambda r: httpx.Response(200, content=b'x' * (sources.MAX_BYTES + 1)), True)
            private.assert_not_called()
        with self.assertRaises(HTTPException) as caught:
            await self.fetch(url, lambda r: httpx.Response(200, content=b'<html>login', headers={'content-type': 'text/html'}))
        self.assertIn('HTML', caught.exception.detail)

    async def test_login_redirect_private_fallback(self):
        seen = []

        def handler(request):
            seen.append(str(request.url))
            return httpx.Response(302, headers={
                'location': 'https://accounts.google.com/ServiceLogin?continue=example',
            })

        with patch.dict(os.environ, {'WORDS_GOG_PRIVATE': '1'}), patch.object(
            sources, '_read_private', return_value=b'private'
        ) as private:
            result = await self.fetch('https://docs.google.com/spreadsheets/d/abc/edit', handler, True)
            self.assertEqual(result, b'private')
            private.assert_called_once_with('abc', True)
        self.assertEqual(seen, ['https://docs.google.com/spreadsheets/d/abc/export?format=xlsx'])

    async def test_private_argv(self):
        for spreadsheet in (True, False):
            def run(argv, **kwargs):
                self.assertEqual(argv[:6], ['/usr/local/bin/gog', 'drive', 'download', 'abc', '--no-input', '--out'])
                self.assertEqual(argv[7:], ['--format=xlsx'] if spreadsheet else [])
                self.assertIs(kwargs['shell'], False)
                self.assertEqual(kwargs['timeout'], 20)
                Path(argv[6]).write_bytes(b'private bytes')
            with patch.object(sources.shutil, 'which', return_value='/usr/local/bin/gog'), patch.object(
                sources.subprocess, 'run', side_effect=run
            ):
                self.assertEqual(sources._read_private('abc', spreadsheet), b'private bytes')
        with patch.object(sources.shutil, 'which', return_value=None), self.assertRaises(HTTPException):
            sources._read_private('abc', True)
        with patch.object(sources.shutil, 'which', return_value='gog'), patch.object(
            sources.subprocess, 'run', side_effect=subprocess.TimeoutExpired('gog', 20)
        ), self.assertRaises(HTTPException):
            sources._read_private('abc', True)


if __name__ == '__main__':
    unittest.main()
