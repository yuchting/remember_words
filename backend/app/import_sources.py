"""XLSX 解析与 Google Drive 下载；调用者负责家长权限及下载后的解析。"""

import asyncio
from io import BytesIO
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from urllib.parse import parse_qs, urljoin, urlsplit
from xml.etree import ElementTree
from zipfile import ZipFile

import httpx
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

MAX_BYTES = 2 * 1024 * 1024
MAX_EXPANDED = 20 * 1024 * 1024
GOOGLE_HOSTS = {'docs.google.com', 'drive.google.com', 'drive.usercontent.google.com'}


def _error(message):
    return HTTPException(422, message)


def _check_archive(data):
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED:
            raise _error('XLSX 展开大小不能超过 20 MiB，文件数不能超过 2000')
        if len({entry.filename for entry in entries}) != len(entries):
            raise _error('XLSX 包含重复 ZIP 文件名')
        total = 0
        for entry in entries:
            with archive.open(entry) as source:
                body = source.read(MAX_EXPANDED + 1)
            total += len(body)
            if total > MAX_EXPANDED:
                raise _error('XLSX 展开大小不能超过 20 MiB')
            if entry.filename.endswith('.xml'):
                markup = body.replace(b'\x00', b'').upper()
                if b'<!DOCTYPE' in markup or b'<!ENTITY' in markup:
                    raise _error('XLSX 不允许 XML DTD 或实体声明')
            if not re.fullmatch(r'xl/worksheets/[^/]+\.xml', entry.filename):
                continue
            # 不信任 dimension 元数据；检查实际坐标，防止越界单元格被截断。
            rows = 0
            for _, node in ElementTree.iterparse(BytesIO(body), events=('end',)):
                tag = node.tag.rsplit('}', 1)[-1]
                if tag == 'row':
                    rows += 1
                    number = int(node.get('r', rows))
                    if rows > 5001 or not 1 <= number <= 5001:
                        raise _error(f'{entry.filename}：行数不能超过 5001（含表头）')
                elif tag == 'c':
                    coordinate = node.get('r', '')
                    match = re.fullmatch(r'([A-Za-z]+)([1-9][0-9]*)', coordinate)
                    if not match:
                        raise _error(f'{entry.filename}：单元格坐标无效')
                    if column_index_from_string(match[1]) > 100 or int(match[2]) > 5001:
                        raise _error(f'{entry.filename}：最多 5001 行、100 列')
                node.clear()


def parse_xlsx(data: bytes) -> dict:
    """逐页解析；无效页返回跳过原因，不阻止其他有效页。"""
    if not isinstance(data, bytes):
        raise _error('XLSX 数据必须是 bytes')
    if len(data) > MAX_BYTES:
        raise _error('XLSX 文件不能超过 2 MiB')
    workbook = None
    try:
        _check_archive(data)
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheets = []
        skipped = []
        for sheet in workbook.worksheets:
            try:
                items = _parse_sheet(sheet)
            except HTTPException as exc:
                skipped.append({'name': sheet.title, 'reason': str(exc.detail)})
                continue
            sheets.append({'name': sheet.title, 'items': items})
        if not sheets:
            reasons = '；'.join(f'{item["name"]}：{item["reason"]}' for item in skipped)
            raise _error('没有可导入的工作表；' + reasons)
        return {'sheets': sheets, 'skipped': skipped}
    except HTTPException:
        raise
    except Exception as exc:
        raise _error('XLSX 无效、损坏或包含不支持的数据') from exc
    finally:
        if workbook is not None:
            workbook.close()


def _parse_sheet(sheet):
    sheet.reset_dimensions()
    headers = None
    items = []
    for number, row in enumerate(sheet.iter_rows(), 1):
        if number > 5001 or len(row) > 100:
            raise _error(f'工作表「{sheet.title}」：最多 5001 行、100 列')
        values = [cell.value for cell in row]
        if all(value is None or isinstance(value, str) and not value.strip() for value in values):
            continue
        if headers is None:
            if number != 1:
                raise _error(f'工作表「{sheet.title}」：第一行必须是表头')
            headers = {}
            for index, value in enumerate(values):
                if value is None or isinstance(value, str) and not value.strip():
                    continue
                if not isinstance(value, str):
                    raise _error(f'工作表「{sheet.title}」：表头必须是文本')
                name = value.strip().casefold()
                name = 'word' if name == 'words' else name
                if name in headers:
                    raise _error(f'工作表「{sheet.title}」：重复表头 {name}')
                headers[name] = index
            missing = {'word', 'meaning'} - headers.keys()
            if missing:
                raise _error(f'工作表「{sheet.title}」：缺少表头 {", ".join(sorted(missing))}')
            continue
        item = {}
        for field in ('word', 'meaning', 'example', 'example_translation'):
            index = headers.get(field)
            cell = row[index] if index is not None and index < len(row) else None
            value = cell.value if cell is not None else None
            if cell is not None and cell.data_type == 'e':
                raise _error(f'工作表「{sheet.title}」第 {number} 行：{field} 是 Excel 错误值')
            if value is None:
                value = ''
            if not isinstance(value, str):
                raise _error(f'工作表「{sheet.title}」第 {number} 行：{field} 必须是文本')
            item[field] = value.strip()
            if field in ('word', 'meaning') and not item[field]:
                raise _error(f'工作表「{sheet.title}」第 {number} 行：{field} 不能为空')
        items.append(item)
    if not items:
        raise _error("空表或没有可导入的单词")
    return items


def _safe_google_url(url, download=False):
    if not isinstance(url, str) or re.search(r'[\s\x00-\x1f\x7f\\]', url):
        raise _error('只允许 HTTPS Google Sheets / Drive 分享链接')
    try:
        parsed = urlsplit(url)
        sheets_download = download and re.fullmatch(
            r'doc-[a-z0-9-]+-sheets\.googleusercontent\.com', parsed.netloc
        )
        if parsed.scheme != 'https' or not (parsed.netloc in GOOGLE_HOSTS or sheets_download):
            raise ValueError
    except ValueError as exc:
        raise _error('只允许 HTTPS 且主机精确匹配的 Google 链接') from exc
    return parsed


def _drive_target(url):
    parsed = _safe_google_url(url)
    identifier = r'([A-Za-z0-9_-]+)'
    if parsed.netloc == 'docs.google.com':
        match = re.fullmatch(r'/spreadsheets/d/' + identifier + r'(?:/(?:edit|view|preview|copy))?/?', parsed.path)
        if match:
            file_id = match[1]
            return file_id, True, f'https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx'
    elif parsed.netloc == 'drive.google.com':
        match = re.fullmatch(r'/file/d/' + identifier + r'(?:/(?:view|edit|preview))?/?', parsed.path)
        ids = parse_qs(parsed.query).get('id', [])
        file_id = match[1] if match else ids[0] if parsed.path == '/open' and len(ids) == 1 else ''
        if re.fullmatch(identifier, file_id):
            return file_id, False, f'https://drive.google.com/uc?export=download&id={file_id}'
    raise _error('链接必须是 Google Sheets /spreadsheets/d/ID 或 Drive /file/d/ID、/open?id=ID')


class _PublicUnavailable(Exception):
    pass


def _check_content(data, content_type=''):
    if len(data) > MAX_BYTES:
        raise _error('下载文件不能超过 2 MiB')
    prefix = data[:1024].lstrip(b'\xef\xbb\xbf \t\r\n').lower()
    if 'html' in content_type.lower() or prefix.startswith((b'<!doctype html', b'<html', b'<head', b'<body')):
        raise _PublicUnavailable('Google 返回 HTML，文件可能未公开分享')
    if not data:
        raise _error('Google 返回空文件')
    return data


async def _read_public(url):
    try:
        async with asyncio.timeout(10):
            async with httpx.AsyncClient(timeout=10, follow_redirects=False, trust_env=False) as client:
                for _ in range(6):
                    _safe_google_url(url, download=True)
                    async with client.stream('GET', url) as response:
                        if response.status_code in (301, 302, 303, 307, 308):
                            location = response.headers.get('location')
                            if not location:
                                raise _PublicUnavailable('Google 重定向缺少 Location')
                            url = urljoin(url, location)
                            login = urlsplit(url)
                            if login.scheme == 'https' and login.netloc == 'accounts.google.com':
                                raise _PublicUnavailable('Google 要求登录，文件未公开访问')
                            _safe_google_url(url, download=True)
                            continue
                        if response.status_code != 200:
                            raise _PublicUnavailable(f'Google 公开下载失败（HTTP {response.status_code}）')
                        size = response.headers.get('content-length', '')
                        if size.isdigit() and int(size) > MAX_BYTES:
                            raise _error('下载文件不能超过 2 MiB')
                        data = bytearray()
                        async for chunk in response.aiter_bytes(chunk_size=65536):
                            data.extend(chunk)
                            if len(data) > MAX_BYTES:
                                raise _error('下载文件不能超过 2 MiB')
                        return _check_content(bytes(data), response.headers.get('content-type', ''))
                raise _error('Google 重定向次数超过 5 次')
    except (httpx.HTTPError, TimeoutError) as exc:
        raise _PublicUnavailable('Google 公开下载网络失败或超时') from exc


def _read_private(file_id, spreadsheet):
    executable = shutil.which('gog')
    if not executable:
        raise _error('私有下载需要安装并授权 gog')
    # 已由本机 gog --help、gog drive download --help 核实。
    export_flags = ['--format=xlsx'] if spreadsheet else []
    try:
        with tempfile.TemporaryDirectory(prefix='words-gog-') as directory:
            output = Path(directory) / 'download.xlsx'
            subprocess.run(
                [executable, 'drive', 'download', file_id, '--no-input', '--out', output.as_posix(), *export_flags],
                shell=False, timeout=20, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            )
            with output.open('rb') as source:
                return _check_content(source.read(MAX_BYTES + 1))
    except _PublicUnavailable as exc:
        raise _error('gog 返回 HTML，无法下载私有文件') from exc
    except (OSError, subprocess.SubprocessError) as exc:
        raise _error('gog 私有下载失败，请检查安装、授权、文件权限或超时') from exc


async def read_drive(url: str, allow_private: bool = False) -> bytes:
    """公开下载失败才回退 gog；家长权限由调用者控制，格式错误不回退。"""
    file_id, spreadsheet, target = _drive_target(url)
    try:
        return await _read_public(target)
    except _PublicUnavailable as exc:
        if allow_private is True and os.environ.get('WORDS_GOG_PRIVATE') == '1':
            return await asyncio.to_thread(_read_private, file_id, spreadsheet)
        raise _error(f'{exc}；请公开分享，或由家长启用私有下载') from exc
